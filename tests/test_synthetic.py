r"""Проверка физики на искусственной съёмке: капли известного радиуса летят с известной скоростью.

Рисуются кадры (светлый фон, тёмные круглые капли), таблица записей и настройки — во временной
папке. Затем по очереди выполняются st1, st2 и st4 (функции process_record из main/) и проверяется,
что найдены все капли, их радиусы и скорость, а средние в области st4 в мм, м/с и мкм совпадают
с заданными. Отдельно — что окно просмотра st1 (tools/st1_viewer.py) показывает все виды и что
приложение (app/) открывает записи, показывает все виды, заполняет и сохраняет таблицу и настройки,
запускает стадии в фоновых процессах (с тем же результатом, что скрипты main/) и показывает результаты.
От конкретного эксперимента и от MATLAB проверка не зависит.

Запуск:  .venv\Scripts\python.exe -B -m pytest tests\test_synthetic.py
"""

import importlib.util
import os
import shutil
import time
from pathlib import Path
from types import ModuleType, SimpleNamespace

import matplotlib
import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

from lib.common import files
from lib.common.records import read_records

matplotlib.use("Agg")  # окно просмотра рисуется в память, на экран не выводится
ROOT = Path(__file__).resolve().parents[1]

WIDTH, HEIGHT = 640, 384
N_FRAMES = 14  # кадров в диапазоне записи; перед ними — кадр фона
VELOCITY = (2.5, -1.5)  # px/кадр (y вниз)
BACKGROUND, DROPLET = 200.0, 60.0  # яркость фона и капель
SCALE, FPS = 0.05, 4000.0  # мм/px, кадров/с
CX, CY = 0.0, 400.0  # начало координат st4, px

# капли: (x, y на первом кадре диапазона, радиус), px — три ряда по пять
DROPLETS = [(x, y, r)
            for y, radii in ((110, (4, 5, 6, 7, 8)), (200, (8, 7, 6, 5, 4)), (290, (5, 6, 7, 8, 6)))
            for x, r in zip((90, 200, 310, 420, 530), radii)]

THRESHOLDS = SimpleNamespace(
    RIM_GRAY_RANGE=(-110, 0), RIM_BINARIZE_THRESHOLD=0.15, RIM_CLOSE_RADIUS=5, RIM_DILATE_RADIUS=5,
    DROPLET_GRAY_RANGE=(0.05, 1.0), CANNY_THRESHOLD=0.1, DROPLET_OPEN_RADIUS=1, DROPLET_MIN_PIXELS=5,
    DROPLET_CLOSE_RADIUS=5,
    PIV_WINDOW=128, PIV_STEP=64, PIV_MAX_DISPLACEMENT=30, PIV_MIN_PEAK_RATIO=1.0,
    PIV_FILTER_MIN_PEAK_RATIO=1.2, PIV_FILTER_MIN_DISPLACEMENT=1.0, PIV_FILTER_MEDIAN_SIZE=3,
    PTV_SEARCH_RADIUS=20, PTV_MAX_RADIUS_RATIO=2.0, PTV_MAX_FRAME_GAP=3, PTV_MAX_DEVIATION=0.05,
    PTV_MIN_STRAIGHTNESS=0.995, PTV_PIV_WEIGHT=0.0, PTV_POLY_POINTS=16,
    TRACKS_PLOT_MIN_POINTS=5, AREA_PLOT_ARROW_SCALE=0.62, AREA_PLOT_Y_LIMITS_MM=(5, 60),
    PIV_PLOT_ARROW_SCALE=10, PIV_PLOT_SPEED_MAX=7,
)


def _frame(k: int | None) -> np.ndarray:
    """Кадр k (с 1) диапазона; None — фон без капель. Края капель сглажены (4×4 подвыборки)."""
    sub = 4
    yy, xx = np.mgrid[0 : HEIGHT * sub, 0 : WIDTH * sub]
    x = (xx + 0.5) / sub + 0.5  # координаты пикселей с 1, как в MATLAB
    y = (yy + 0.5) / sub + 0.5
    inside = np.zeros(x.shape, dtype=bool)
    if k is not None:
        for x0, y0, r in DROPLETS:
            cx, cy = x0 + VELOCITY[0] * (k - 1), y0 + VELOCITY[1] * (k - 1)
            inside |= (x - cx) ** 2 + (y - cy) ** 2 <= r**2
    coverage = inside.reshape(HEIGHT, sub, WIDTH, sub).mean(axis=(1, 3))
    return BACKGROUND + (DROPLET - BACKGROUND) * coverage


def _script(folder: str, name: str) -> ModuleType:
    spec = importlib.util.spec_from_file_location(name, ROOT / folder / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def run(tmp_path_factory) -> SimpleNamespace:
    """Искусственная запись и результаты всех трёх стадий."""
    root = tmp_path_factory.mktemp("synthetic")
    images = root / "images" / "rec1"
    images.mkdir(parents=True)
    Image.fromarray(np.round(_frame(None)).astype(np.uint8)).save(images / "rec1_0001.png")
    for k in range(1, N_FRAMES + 1):
        Image.fromarray(np.round(_frame(k)).astype(np.uint8)).save(images / f"rec1_{k + 1:04d}.png")

    wb = Workbook()
    ws = wb.active
    ws.append(["cond_num", "rec_name", "images_path_side", "start_frame", "fin_frame", "rim", "cx", "cy"])
    ws.append(["описание"] * 8)
    ws.append([1, "rec1", "rec1", 2, N_FRAMES + 1, 5000, CX, CY])
    wb.save(root / "table.xlsx")

    settings = SimpleNamespace(
        IMAGES_DIR=str(root / "images"), TABLE_FILE=str(root / "table.xlsx"), OUTPUT_DIR=str(root / "out"),
        CONDITIONS=[], RECORDS=[], SCALE_MM_PER_PX=SCALE, FPS=FPS,
        SURFACE_ROW=HEIGHT + 1, BACKGROUND_FRAME=1, ERASE_REGIONS=[],  # воды в кадре нет; фон — первый файл
        AREA_X=(0, WIDTH), AREA_HEIGHTS=[60], AREA_LOW_EDGES=[150, 240, 330],
    )
    (record,) = read_records(Path(settings.TABLE_FILE))
    for stage in ("st1_droplets_separation", "st2_ptv", "st4_droplets_in_area"):
        _script("main", stage).process_record(record, settings, THRESHOLDS)
    out = Path(settings.OUTPUT_DIR)
    return SimpleNamespace(settings=settings, record=record, out=out,
                           st1=record.st1_dir(out), st2=record.st2_dir(out), st4=record.st4_dir(out))


def _true_centers(k: int) -> np.ndarray:
    return np.array([[x + VELOCITY[0] * (k - 1), y + VELOCITY[1] * (k - 1)] for x, y, _ in DROPLETS])


def test_st1_finds_every_droplet_with_its_position_and_radius(run):
    centers = files.load_cell(run.st1 / "particles.mat", "centers_arr")
    radii = files.load_cell(run.st1 / "particles.mat", "radii_arr")
    true_r = np.array([r for _, _, r in DROPLETS], dtype=float)
    for k in range(1, N_FRAMES + 1):
        found, truth = centers[k - 1].reshape(-1, 2), _true_centers(k)
        distance = np.hypot(*(found[:, None, :] - truth[None, :, :]).transpose(2, 0, 1))
        nearest = distance.argmin(axis=0)  # ближайшая найденная капля для каждой настоящей
        assert distance.min(axis=0).max() < 0.5, f"кадр {k}: капля найдена дальше 0.5 px от настоящей"
        assert len(found) == len(truth), f"кадр {k}: найдено {len(found)} капель вместо {len(truth)}"
        # Метод MATLAB завышает радиус на 0.2–0.6 px — вопрос к специалисту Q11 (QUESTIONS.md), отсюда допуск 0.7 px
        np.testing.assert_allclose(radii[k - 1].ravel()[nearest], true_r, atol=0.7, err_msg=f"кадр {k}")


def test_st1_piv_velocity_at_droplets(run):
    xd = files.load_cell(run.st1 / "particles.mat", "xd_q_arr")
    yd = files.load_cell(run.st1 / "particles.mat", "yd_q_arr")
    vx = np.concatenate([a.ravel() for a in xd])
    vy = np.concatenate([a.ravel() for a in yd])
    ok = ~np.isnan(vx)
    assert ok.mean() > 0.9, "скорость PIV не определена у многих капель"
    np.testing.assert_allclose(np.median(vx[ok]), VELOCITY[0], atol=0.1)
    np.testing.assert_allclose(np.median(vy[ok]), VELOCITY[1], atol=0.1)


def test_st2_one_long_track_per_droplet(run):
    tracks = files.load_cell(run.st2 / "ptv_unfiltered.mat", "p_arr")
    long_tracks = [t for t in tracks if t.shape[0] >= N_FRAMES - 3]
    assert len(long_tracks) == len(DROPLETS), f"длинных траекторий {len(long_tracks)} вместо {len(DROPLETS)}"
    for t in long_tracks:
        frames, x, y = t[:, 0], t[:, 3], t[:, 4]
        np.testing.assert_allclose(np.diff(x) / np.diff(frames), VELOCITY[0], atol=0.3)
        np.testing.assert_allclose(np.diff(y) / np.diff(frames), VELOCITY[1], atol=0.3)


def test_st4_means_in_physical_units(run):
    stat = files.load_dat(run.st4 / "_all_stat.dat")
    assert stat.shape[0] == 3, "в каждой из трёх областей должны быть траектории"
    vx_true = VELOCITY[0] * SCALE / 1000 * FPS  # м/с
    vy_true = -VELOCITY[1] * SCALE / 1000 * FPS  # м/с, ось y вверх
    np.testing.assert_allclose(stat[:, 3], vx_true, rtol=0.03)
    np.testing.assert_allclose(stat[:, 4], vy_true, rtol=0.03)
    for (low, n_tracks, r_mean) in zip(stat[:, 1], stat[:, 2], stat[:, 5]):
        assert n_tracks >= 5, f"область с нижним краем {low} мм: всего {n_tracks} траекторий"
        assert 4 * SCALE * 1000 * 0.9 < r_mean < 8 * SCALE * 1000 * 1.1  # мкм, в пределах размеров капель


def _viewer(run):
    return _script("tools", "st1_viewer").Viewer(run.settings, THRESHOLDS, workers=0)  # без фоновых процессов


def test_st1_viewer_shows_every_view_with_matching_caption(run):
    viewer = _viewer(run)
    assert "кадр 1 из" in viewer.ax.get_title() and "капель: Python 15" in viewer.ax.get_title()
    viewer.ax.set_xlim(100, 300)  # увеличение сохраняется при листании и смене вида
    for key in ("1", "2", "3", "4", "5", "m", "3"):
        viewer.press(key)
        assert "ошибка" not in viewer.ax.get_title(), f"клавиша {key}: {viewer.ax.get_title()}"
    viewer.press("right")
    viewer.release("right")
    assert viewer.k == 2 and "кадр 2 из" in viewer.ax.get_title()
    assert viewer.ax.get_xlim() == (100, 300)
    viewer.press("end")
    viewer.press("5")
    assert viewer.k == N_FRAMES and f"кадр {N_FRAMES} из" in viewer.ax.get_title()
    viewer.close()


def test_st1_viewer_slideshow_by_space_and_by_holding_arrow(run):
    viewer = _viewer(run)
    viewer.press(" ")  # пуск: кадр на каждый тик таймера
    for _ in range(5):
        viewer.tick()
    assert viewer.k == 6 and "кадр 6 из" in viewer.ax.get_title()
    viewer.press(" ")  # стоп
    viewer.tick()
    assert viewer.k == 6
    viewer.press("left")  # шаг назад сразу ...
    viewer.held_since -= 1  # ... стрелку держат уже секунду — слайдшоу назад
    viewer.tick()
    viewer.tick()
    assert viewer.k == 3
    viewer.release("left")
    viewer.tick()
    assert viewer.k == 3
    viewer.close()


def test_st1_viewer_computes_frames_exactly_like_st1(run):
    from lib.st1.preview import RecordPreview

    preview = RecordPreview(run.record, run.settings, THRESHOLDS)
    centers, radii, xd, yd = files.load_particles(run.st1)
    for k in (1, 7, N_FRAMES):  # последний кадр: скорость по PIV последней пары, как в st1
        result = preview.result(k)
        np.testing.assert_array_equal(result.centers, centers[k - 1])
        np.testing.assert_array_equal(result.radii, radii[k - 1])
        np.testing.assert_array_equal(result.xd, xd[k - 1])
        np.testing.assert_array_equal(result.yd, yd[k - 1])


def test_app_opens_records_shows_every_view_and_saves_the_table(run, tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")  # окно без экрана
    from PySide6.QtWidgets import QApplication

    from app.records_page import VIEW_KEYS
    from app.window import MainWindow

    application = QApplication.instance() or QApplication([])  # noqa: F841 — нужно окну
    # копия таблицы с ещё одной, незаполненной записью rec2 (её кадры — копия кадров rec1)
    table = tmp_path / "table.xlsx"
    workbook = load_workbook(run.settings.TABLE_FILE)
    workbook.active.append([1, "rec2", "rec2"])
    workbook.save(table)
    for name in ("rec1", "rec2"):
        shutil.copytree(Path(run.settings.IMAGES_DIR) / "rec1", tmp_path / "images" / name)
    settings = SimpleNamespace(**{**vars(run.settings), "TABLE_FILE": str(table),
                                  "IMAGES_DIR": str(tmp_path / "images"), "OUTPUT_DIR": str(tmp_path / "out")})
    window = MainWindow(settings, THRESHOLDS, workers=0)  # кадры считаются сразу, без фоновых процессов
    page = window.records

    assert page.list.rowCount() == 2 and page.list.item(0, 1).text() == "✓"
    assert "кадр 2 из 15" in page.title.text()  # rec1 открыта на начале диапазона (файл 2)
    for view in VIEW_KEYS.values():
        page.set_view(view)
        assert "шибка" not in page.info.text(), f"{view}: {page.info.text()}"
    page.set_view("droplets")
    assert "капель: Python 15" in page.info.text()
    page.toggle_slideshow()
    for _ in range(3):
        page._slideshow_step()  # шаг таймера слайдшоу
    page.toggle_slideshow()
    assert page.number == 5

    page.switch_record(1)  # rec2: нет диапазона, rim и центра
    assert page.i == 1 and "кадр 1 из 15" in page.title.text() and "нужен rim" in page.info.text()
    page.fields["rim"].hint.click()  # подсказка «← 5000» — rim предыдущей заполненной записи
    page.recompute()  # то же делает таймер после правки rim
    page.go(3)  # файл 1 — фон, капли — с файла 2
    assert "капель: Python 15" in page.info.text()
    page.start_button.click()
    page.go(12)
    page.fin_button.click()
    page.center_button.click()
    page.view.clicked.emit(100.4, 200.6)  # щелчок по кадру в режиме «центр затравки»
    assert not page.center_button.isChecked() and window.windowTitle().endswith("*")

    assert window.save() and not page.has_changes
    rec1, rec2 = read_records(table)
    assert (rec2.start_frame, rec2.fin_frame, rec2.rim, rec2.cx, rec2.cy) == (3, 12, 5000, 100, 201)
    assert rec1 == read_records(Path(run.settings.TABLE_FILE))[0]  # rec1 не тронута
    assert list((tmp_path / "out" / "table_backups").glob("table *.xlsx"))  # копия таблицы до записи
    window.close()


def _module_file(path: Path, values: dict) -> ModuleType:
    """Файл настроек (ИМЯ = значение  # комментарий) и модуль из него — как settings.py, thresholds.py."""
    path.write_text("".join(f"{name} = {value!r}  # {name.lower()}\n" for name, value in values.items()),
                    encoding="utf-8")
    spec = importlib.util.spec_from_file_location(path.stem, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_app_settings_panel_recomputes_frame_and_saves_files(run, tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")  # окно без экрана
    from PySide6.QtWidgets import QApplication

    from app.window import MainWindow

    application = QApplication.instance() or QApplication([])  # noqa: F841 — нужно окну
    settings = _module_file(tmp_path / "settings.py", {**vars(run.settings), "OUTPUT_DIR": str(tmp_path / "out")})
    thresholds = _module_file(tmp_path / "thresholds.py", vars(THRESHOLDS))
    window = MainWindow(settings, thresholds, workers=0)
    page, panel = window.records, window.records.params
    page.go(3)
    page.set_view("droplets")
    assert "капель: Python 15" in page.info.text()

    panel.set_text("thresholds", "RIM_GRAY_RANGE", "(-300, -200)")  # всё светлее границы — капель нет
    page.apply_parameters()  # то же делает таймер после правки
    assert "капель: Python 0" in page.info.text() and window.windowTitle().endswith("*")
    field = panel.fields[("thresholds", "RIM_BINARIZE_THRESHOLD")]
    field.sliders[0].setValue(800)  # ползунок 0…0.3 (вдвое больше значения 0.15)
    assert panel.files["thresholds"].edits["RIM_BINARIZE_THRESHOLD"] == field.edit.text() == "0.24"
    panel.set_text("thresholds", "RIM_CLOSE_RADIUS", "(1, 2)")  # не число — красным, правки нет
    assert "ffc9c9" in panel.fields[("thresholds", "RIM_CLOSE_RADIUS")].edit.styleSheet()
    assert "RIM_CLOSE_RADIUS" not in panel.files["thresholds"].edits
    panel.set_text("thresholds", "RIM_BINARIZE_THRESHOLD", "0.15")  # как в файле — правка снимается
    assert "RIM_BINARIZE_THRESHOLD" not in panel.files["thresholds"].edits

    assert window.save() and not panel.has_changes and not window.windowTitle().endswith("*")
    text = (tmp_path / "thresholds.py").read_text(encoding="utf-8")
    assert "RIM_GRAY_RANGE = (-300, -200)  # rim_gray_range\n" in text and "RIM_BINARIZE_THRESHOLD = 0.15  #" in text
    assert thresholds.RIM_GRAY_RANGE == (-300, -200)  # модуль в памяти — как файл

    panel.set_text("thresholds", "RIM_GRAY_RANGE", "(-110, 0)")
    panel.discard()  # «Вернуть как в файлах»
    page.apply_parameters()
    assert not panel.has_changes and "капель: Python 0" in page.info.text()
    window.close()


def test_app_runs_stages_like_main_scripts_and_shows_results(run, tmp_path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")  # окно без экрана
    from PySide6.QtWidgets import QApplication

    from app.window import MainWindow

    application = QApplication.instance() or QApplication([])
    settings = _module_file(tmp_path / "settings.py", {**vars(run.settings), "OUTPUT_DIR": str(tmp_path / "out")})
    thresholds = _module_file(tmp_path / "thresholds.py", vars(THRESHOLDS))
    window = MainWindow(settings, thresholds, workers=0)
    jobs, results = window.jobs, window.results

    def wait_for_jobs(seconds: float = 300) -> None:
        start = time.monotonic()
        while jobs.running and time.monotonic() - start < seconds:
            application.processEvents()  # таймер вкладки забирает сообщения процессов
            time.sleep(0.05)

    window.tabs.setCurrentWidget(jobs)
    jobs.set_checked([0])
    jobs.start()  # st1 → st2 → st4 в отдельном процессе
    wait_for_jobs()
    assert jobs.list.item(0, 5).text() == "st4 ✓" and jobs.progress.text() == "готово 1 из 1"
    assert "кадр 14 из 14" in jobs.log.toPlainText()  # ход работы st1 дошёл до окна

    out = Path(settings.OUTPUT_DIR)  # результаты — те же, что у скриптов main/ (фикстура run)
    for name in ("centers_arr", "radii_arr", "xd_q_arr"):
        ours, theirs = (files.load_cell(folder / "particles.mat", name) for folder in (run.record.st1_dir(out), run.st1))
        assert all(np.array_equal(a, b, equal_nan=True) for a, b in zip(ours, theirs))
    assert (run.record.st4_dir(out) / "_all_stat.dat").read_bytes() == (run.st4 / "_all_stat.dat").read_bytes()
    assert window.records.list.item(0, 4).text() == "✓"  # отметка «есть st4» на «Записях»

    window.tabs.setCurrentWidget(results)
    assert results.tracks_info.text().startswith("траекторий") and results.summary.rowCount() == 3
    results.summary.selectRow(1)
    assert results.area_table.rowCount() >= 5 and "траекторий" in results.area_info.text()

    jobs.start()
    jobs.stop()  # «Остановить» сразу после запуска
    assert jobs.progress.text() == "готово 0 из 1, остановлено 1" and not jobs.running
    window.close()


def test_job_runner_reports_error_and_stops_the_chain(run, tmp_path):
    from dataclasses import replace

    from lib.common.jobs import Job, JobRunner

    settings = _module_file(tmp_path / "settings.py", {**vars(run.settings), "OUTPUT_DIR": str(tmp_path / "out")})
    _module_file(tmp_path / "thresholds.py", vars(THRESHOLDS))
    missing = replace(run.record, path="нет такой записи")
    runner = JobRunner(2)
    runner.submit([Job(1, missing, ("st1", "st2"), settings.__file__, str(tmp_path / "thresholds.py"))])
    messages, start = [], time.monotonic()
    while (runner.busy() or not messages) and time.monotonic() - start < 120:
        messages += runner.poll()
        time.sleep(0.05)
    kinds = [(kind, stage) for _, kind, stage, _ in messages if kind != "log"]
    assert kinds == [("start", "st1"), ("error", "st1")]  # st2 после ошибки st1 не запускается
    assert "Нет папки с кадрами" in messages[-1][3]
