r"""Проверки отдельных функций на маленьких примерах с известным ответом. Данные не нужны.

Запуск:  .venv\Scripts\python.exe -B -m pytest
"""

import base64
import csv
import lzma
import math
import xml.etree.ElementTree as ET
from pathlib import Path

import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import PatternFill

from lib.analysis.dataset import DataTable
from lib.analysis.formulas import computed_columns, evaluate, fit, names_in, parameters_of
from lib.analysis.labplot import write_csv, write_lml
from lib.analysis.stats import Condition, binned, histogram, select, split, window_edges
from lib.common.config_edit import SettingsFile, current_experiment, set_current_experiment, value_text
from lib.common.files import last_number_in_name, save_dat
from lib.common.records import read_records
from lib.common.table_edit import EditableTable, TableSaveError
from lib.st1 import image_ops as ops
from lib.st1.piv import PivGrid, PivResult, filter_piv, median_ignoring_nan
from lib.st2.tracking import PolyFit
from lib.st4.areas import px_per_frame_to_m_per_s, px_to_mm, segments_intersect


# ------------------------------------------------------------------ файлы и таблица


def test_frames_are_sorted_by_last_number_in_name():
    assert last_number_in_name("0501") == 501
    assert last_number_in_name("asd_0502.jpg") == 502
    assert last_number_in_name("asdf1231asd2_0501.jpg") == 501
    assert math.isnan(last_number_in_name("sadfasdf.jpg"))


def test_dat_file_is_byte_identical_to_matlab(tmp_path):
    path = tmp_path / "a.dat"
    save_dat(path, np.array([[0.6263, -2.4630237, 58.74694]]))
    assert path.read_bytes() == b"   6.2630000e-01  -2.4630237e+00   5.8746940e+01\r\n"
    save_dat(path, 0.0)
    assert path.read_bytes() == b"   0.0000000e+00\r\n"


def _table(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    ws.append(["описание"] * len(headers))  # вторая строка — описания, пропускается
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def test_records_accept_both_column_spellings_and_both_path_layouts(tmp_path):
    # таблица как v4_Zot_new: rim_new, cx; путь уже оканчивается именем записи
    t1 = _table(tmp_path / "a.xlsx",
                ["cond_num", "rec_name", "images_path_side", "start_frame", "fin_frame", "rim_new", "cx", "cy"],
                [[3, "F40_1", r"1 F40 z1 side\F40_1", 10, 20, 400, 235, 908]])
    (r1,) = read_records(t1)
    assert (r1.path, r1.rim, r1.cx, r1.cy) == ("1 F40 z1 side/F40_1", 400, 235, 908)

    # таблица как clear water 3: rim, c_x; к пути добавляется имя записи; запись без кадров
    t2 = _table(tmp_path / "b.xlsx",
                ["cond_num", "rec_name", "images_path_side", "start_frame", "fin_frame", "rim", "c_x", "c_y"],
                [[1, "0201", r"water\140", 80, 162, 200, 28, 883],
                 [1, "0205", r"water\140", None, None, None, None, None]])
    r2, r3 = read_records(t2)
    assert (r2.name, r2.path, r2.rim, r2.cx) == ("0201", "water/140/0201", 200, 28)
    assert r2.has_frames and not r3.has_frames


def test_records_selection_by_condition_and_name(tmp_path):
    t = _table(tmp_path / "a.xlsx", ["cond_num", "rec_name", "images_path_side"],
               [[1, "a", "x"], [2, "b", "x"], [2, "c", "x"]])
    assert [r.name for r in read_records(t, conditions=[2])] == ["b", "c"]
    assert [r.name for r in read_records(t, names=["a", "c"])] == ["a", "c"]
    assert [r.name for r in read_records(t, conditions=[2], names=["a", "c"])] == ["c"]


TABLE_HEADERS = ["cond_num", "rec_name", "images_path_side", "start_frame", "fin_frame", "rim", "c_x", "c_y"]


def test_table_edit_writes_only_changed_cells_and_keeps_formatting(tmp_path):
    path = _table(tmp_path / "t.xlsx", TABLE_HEADERS + ["comment"],
                  [[1, "0201", r"w`", 80, 162, 200, 28, 883, "хорошая"], [1, "0205", r"w`"]])
    workbook = load_workbook(path)  # форматирование, как в настоящей таблице
    sheet = workbook.active
    sheet.conditional_formatting.add("A1:B2", ColorScaleRule(start_type="min", start_color="FF0000",
                                                             end_type="max", end_color="00FF00"))
    sheet["F3"].fill = PatternFill("solid", fgColor="FFFF00")
    sheet.column_dimensions["B"].width = 33
    workbook.save(path)

    table = EditableTable(path)
    assert table.column("rim") == "rim" and table.column("cx") == "c_x"
    table.set(1, "start_frame", table.parse("start_frame", "41"))
    table.set(1, "rim", table.parse("rim", "200,5"))  # запятая — тоже десятичный знак
    table.set(0, "rim", 200.0)  # то же, что в файле, — не правка
    assert table.changes == {(1, "start_frame"): 41, (1, "rim"): 200.5}
    assert table.parse("rec_name", "0207") == "0207"  # текстовая колонка: ведущий ноль не теряется
    with pytest.raises(ValueError):
        table.parse("fin_frame", "abc")  # в числовой колонке — только число

    backup = table.save(tmp_path / "backups")
    assert backup.exists() and not table.changes
    sheet = load_workbook(path).active
    assert (sheet["D4"].value, sheet["F4"].value, sheet["F3"].value, sheet["I3"].value) == (41, 200.5, 200, "хорошая")
    assert sheet["F3"].fill.fgColor.rgb.endswith("FFFF00") and sheet.column_dimensions["B"].width == 33
    assert len(list(sheet.conditional_formatting)) == 1
    assert load_workbook(backup).active["D4"].value is None  # копия — таблица до записи


def test_table_edit_records_and_autofill_from_filled_records(tmp_path):
    path = _table(tmp_path / "t.xlsx", TABLE_HEADERS,
                  [[1, "0201", r"w`", 80, 162, 200, 28, 883], [1, "0202", r"w`", 41, 101, 200, 28, 883],
                   [1, "0203", r"w`", 32, 86, 300, 28, 883], [1, "0204", r"w`"]])
    table = EditableTable(path)
    assert [table.record(i) for i in range(len(table))] == read_records(path)  # те же правила, что у st1–st4
    assert table.suggestion(3, "rim") == (2, 300)  # подсказка — предыдущая заполненная запись
    assert table.suggestion(0, "rim") is None  # поле не пустое
    # «как у предыдущих» — только колонки, одинаковые во всех заполненных записях (rim у них разный)
    assert table.fill_like_previous(3) == ["c_x", "c_y"]
    assert (table.record(3).cx, table.record(3).cy, table.record(3).rim) == (28, 883, None)


def test_table_edit_adds_missing_column_at_first_free_place(tmp_path):
    path = _table(tmp_path / "t.xlsx", TABLE_HEADERS, [[1, "0201", "w"], [1, "0202", "w"]])
    table = EditableTable(path)
    table.add_column("bag_length", "длина бэга, мм")
    table.add_column("bag_frame", "кадр")  # без значений — в файл не попадёт
    table.set(1, "bag_length", 41.25)
    table.save()
    sheet = load_workbook(path).active
    assert [c.value for c in sheet[1]] == TABLE_HEADERS + ["bag_length"]
    assert (sheet["I2"].value, sheet["I3"].value, sheet["I4"].value) == ("длина бэга, мм", None, 41.25)
    assert table.headers[-1] == "bag_length" and not table.new_columns


def test_table_edit_does_not_overwrite_rows_changed_in_excel_or_open_file(tmp_path, monkeypatch):
    path = _table(tmp_path / "t.xlsx", TABLE_HEADERS, [[1, "0201", "w"], [1, "0202", "w"]])
    table = EditableTable(path)
    table.set(1, "rim", 100)
    workbook = load_workbook(path)
    workbook.active["B4"] = "другая"  # тем временем строку переписали в Excel
    workbook.save(path)
    with pytest.raises(TableSaveError, match="правили"):
        table.save()

    table = EditableTable(path)
    table.set(1, "rim", 100)

    def locked(self, data):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(Path, "write_bytes", locked)  # так ведёт себя файл, открытый в Excel
    with pytest.raises(TableSaveError, match="открыта в Excel"):
        table.save()
    assert table.changes  # правки не потеряны


# ------------------------------------------------------------------ файлы настроек для приложения

SETTINGS_TEXT = (
    "# Пороги.\r\n\r\n"
    "# ---------- st1: удаление рима\r\n"
    "RIM_GRAY_RANGE = (-100, 0)  # «кадр минус фон»\r\n"
    "CANNY_THRESHOLD = 0.1  # порог canny\r\n"
    "# Пояснение над строкой\r\n"
    "ERASE_REGIONS = []\r\n"
    "IMAGES_DIR = r\"D:\\ИПФ РАН\\кадры\"  # корень\r\n"
)


def test_settings_file_writes_only_changed_values_and_keeps_comments(tmp_path):
    path = tmp_path / "thresholds.py"
    path.write_bytes(SETTINGS_TEXT.encode("utf-8"))  # концы строк CRLF должны сохраниться
    f = SettingsFile(path)
    rim = f.settings["RIM_GRAY_RANGE"]
    assert (rim.value, rim.comment, rim.section) == ((-100, 0), "«кадр минус фон»", "st1: удаление рима")
    assert f.settings["ERASE_REGIONS"].note == "Пояснение над строкой"
    assert f.settings["IMAGES_DIR"].value == "D:\\ИПФ РАН\\кадры"

    f.set_text("CANNY_THRESHOLD", "0.0000001")  # пишется как набрано, не 1e-07
    f.set_text("ERASE_REGIONS", "[(381, None, 1, 302)]")
    f.set_text("IMAGES_DIR", value_text("D:\\ИПФ РАН\\другие кадры"))
    f.set_text("RIM_GRAY_RANGE", "(-100, 0)")  # как в файле — не правка
    with pytest.raises(ValueError, match="нужно число"):
        f.set_text("CANNY_THRESHOLD", "(1, 2)")
    with pytest.raises(ValueError):
        f.set_text("CANNY_THRESHOLD", "abc")
    assert set(f.edits) == {"CANNY_THRESHOLD", "ERASE_REGIONS", "IMAGES_DIR"}

    assert sorted(f.save()) == ["CANNY_THRESHOLD", "ERASE_REGIONS", "IMAGES_DIR"]
    expected = (SETTINGS_TEXT.replace("0.1  #", "0.0000001  #").replace("= []", "= [(381, None, 1, 302)]")
                .replace("кадры\"", "другие кадры\""))
    assert path.read_bytes() == expected.encode("utf-8")  # всё остальное — байт в байт
    assert f.value("CANNY_THRESHOLD") == 1e-7 and not f.edits

    f.set_text("CANNY_THRESHOLD", "0.2")
    outside = "# дописали в редакторе\r\n".encode("utf-8")
    path.write_bytes(path.read_bytes() + outside)  # файл изменили снаружи
    with pytest.raises(ValueError, match="изменили снаружи"):
        f.save()  # чужая правка не затирается
    assert path.read_bytes().endswith(outside) and f.edits


def test_current_experiment_is_switched_in_current_py(tmp_path):
    path = tmp_path / "current.py"
    path.write_text("# Какой эксперимент.\n#   from experiments.x import settings_2 as settings, thresholds\n"
                    "from experiments.clear_water_3 import settings, thresholds\n", encoding="utf-8")
    assert current_experiment(path) == "clear_water_3"
    set_current_experiment(path, "v4_zot_new")
    assert current_experiment(path) == "v4_zot_new"
    assert path.read_text(encoding="utf-8").startswith("# Какой эксперимент.\n#   from experiments.x")


# ------------------------------------------------------------------ st1


def test_objects_are_numbered_by_columns_like_matlab():
    bw = np.zeros((4, 4), dtype=bool)
    bw[3, 0] = True  # первый при обходе по столбцам
    bw[0, 2] = True  # был бы первым при обходе по строкам
    lab, n = ops.label(bw)
    assert n == 2 and lab[3, 0] == 1 and lab[0, 2] == 2


def test_centroid_is_in_pixels_from_one():
    bw = np.zeros((5, 5), dtype=bool)
    bw[0, 0] = True
    centers, diameters = ops.centroids_and_diameters(bw)
    assert centers.tolist() == [[1.0, 1.0]]
    assert math.isclose(diameters[0], math.sqrt(4 / math.pi))


def test_disk_is_matlab_octagon():
    se = ops.disk(5)
    assert se.shape == (9, 9) and se.sum() == 69  # 81 − 4 угла по 3 пикселя
    yy, xx = np.mgrid[-3:4, -3:4]
    assert np.array_equal(ops.disk(4), np.abs(xx) + np.abs(yy) <= 4)  # strel('disk', 4) MATLAB: 7×7
    assert ops.disk(1).sum() == 5 and ops.disk(2).shape == (5, 5)  # R < 3 — обычный круг


def test_manual_erase_regions_become_background():
    from lib.st1.droplets import RimRemover

    background = np.full((20, 30), 200.0)
    frame = background - 50  # весь кадр темнее фона на 50
    rim = RimRemover((-100, 0), 0.9, 1, 1, surface_row=21, erase_regions=[(5, None, 1, 10)])
    clean, mask = rim.remove_with_mask(frame, background, rim_area=1e9)
    assert (clean[4:, :10] == 1).all() and mask[4:, :10].all()  # строки 5…конец, столбцы 1…10
    assert (clean[:4, :] == 0.5).all() and not mask[:4, :].any() and not mask[:, 10:].any()


def test_median_ignores_nan_and_keeps_nan_where_it_was():
    a = np.array([[1.0, np.nan, 3.0], [4.0, 5.0, 6.0], [np.nan, 8.0, 9.0]])
    m = median_ignoring_nan(a, 3)
    assert math.isnan(m[0, 1]) and math.isnan(m[2, 0])
    assert m[1, 1] == 5.0  # медиана 1, 3, 4, 5, 6, 8, 9


def _piv(status, ratio, xd, yd):
    a = lambda v: np.array([[v]], dtype=float)  # noqa: E731
    return PivResult(a(status), a(xd), a(yd), a(1.0), a(ratio), a(0.0),
                     PivGrid.regular(64, 64, 64), {}, (64, 64), 0.0)


def test_piv_filter_as_in_matlab():
    def kept(status, ratio, xd, yd):
        return not np.isnan(filter_piv(_piv(status, ratio, xd, yd), 1.2, 1.0, 1).xdispl[0, 0])

    assert kept(0, 2.0, 2.0, 0.0)  # хороший вектор
    assert not kept(5, 2.0, 2.0, 0.0)  # плохой status
    assert not kept(0, 1.0, 2.0, 0.0)  # мало отношение пиков
    assert not kept(0, 2.0, 0.5, 0.5)  # почти нет смещения
    # Q1 (вопрос к специалисту): как в MATLAB, при большом yd status и отношение пиков не проверяются
    assert kept(5, 0.0, 0.0, 2.0)


# ------------------------------------------------------------------ st2


def test_polyfit_matches_numpy_and_delta_is_infinite_without_freedom():
    x = np.array([10.0, 9.0, 8.0, 7.0, 6.0])
    y = 0.5 * x**2 - 3 * x + 4 + np.array([0.01, -0.02, 0.0, 0.02, -0.01])
    fit = PolyFit(x, y, 2)
    assert np.allclose(fit.coeffs, np.polyfit(x, y, 2))
    value, delta = fit.value_and_delta(5.0)
    assert math.isclose(value, np.polyval(fit.coeffs, 5.0))
    assert fit.df == 2 and delta > 0
    assert math.isinf(PolyFit(x[:3], y[:3], 2).value_and_delta(5.0)[1])


# ------------------------------------------------------------------ st4


def test_segments_intersect_including_ends_and_not_parallel():
    segments = np.array([[0, 0, 2, 2], [0, 1, 1, 1], [5, 5, 6, 6]], dtype=float)
    border = np.array([[1, 0, 1, 3]], dtype=float)
    assert segments_intersect(segments, border).ravel().tolist() == [True, True, False]
    assert not segments_intersect(np.array([[0, 0, 1, 0]], float), np.array([[0, 1, 1, 1]], float)).any()


def test_units():
    assert px_to_mm(np.array([[120.0, 948.0]]), 0.5, 100.0, 1000.0).tolist() == [[10.0, 26.0]]
    v = px_per_frame_to_m_per_s(np.array([[2.0, 3.0]]), 0.06263, 3990)
    assert np.allclose(v, [[2 * 0.06263 / 1000 * 3990, -3 * 0.06263 / 1000 * 3990]])


# ------------------------------------------------------------------ обработка результатов


def test_selection_windows_histogram_and_binned_statistics():
    table = DataTable({"record": np.array(["a", "a", "b", "b", "b"], dtype=object),
                       "F": np.array([140.0, 140.0, 160.0, 160.0, 160.0]),
                       "r": np.array([1.0, 2.0, 2.5, 3.0, 4.0]), "v": np.array([10.0, 12.0, 5.0, 7.0, 9.0])})
    assert len(select(table, [Condition("F", (160.0,)), Condition("r", low=2.5, high=3.0)])) == 2
    assert len(select(table, [Condition("record", ("a",))])) == 2
    edges = window_edges(table.columns["r"], width=1.0, low=1.0)
    assert edges.tolist() == [1.0, 2.0, 3.0, 4.0]  # до максимума r = 4
    assert histogram(table.columns["r"], edges).y.tolist() == [1, 2, 2]  # 4 — в последнее окно
    density = histogram(table.columns["r"], np.array([0.0, 2.0, 4.0]), "density")
    assert np.allclose(density.y, [1 / 5 / 2, 4 / 5 / 2])  # доля / ширина окна
    s = binned(table.columns["r"], table.columns["v"], edges, "mean_std")
    assert np.allclose(s.y, [10, 8.5, 8]) and np.allclose(s.high - s.y, [0, np.std([12, 5], ddof=1), np.std([7, 9], ddof=1)])
    median = binned(table.columns["r"], table.columns["v"], np.array([0.0, 5.0]), "median")
    assert median.y[0] == 9 and (median.low[0], median.high[0]) == (7, 10)
    assert [label for label, part in split(table, "F")] == ["F = 140", "F = 160"]


def test_labplot_project_and_csv_hold_the_same_numbers(tmp_path):
    columns = {"record": np.array(["a", "b"], dtype=object), "r": np.array([1.5, np.nan]), "v": np.array([2.0, 3.0])}
    write_lml(tmp_path / "p.lml", {"выборка": columns, "кривая": {"x": np.array([0.5]), "y": np.array([7.0])}})
    project = ET.fromstring(lzma.decompress((tmp_path / "p.lml").read_bytes()).decode("utf-8"))
    assert project.tag == "project" and project.get("xmlVersion") == "7"
    sheets = project.findall("child_aspect/spreadsheet")
    assert [sheet.get("name") for sheet in sheets] == ["выборка", "кривая"]
    record, r, v = sheets[0].findall("column")
    assert record.get("mode") == "1" and [row.text for row in record.findall("row")] == ["a", "b"]
    data = np.frombuffer(base64.b64decode(r.find("output_filter").tail), dtype="<f8")  # числа — после фильтров
    assert r.get("mode") == "0" and r.get("designation") == "2" and data[0] == 1.5 and np.isnan(data[1])
    assert record.get("designation") == "1"  # первая колонка — X

    write_csv(tmp_path / "t.csv", columns)
    rows = list(csv.reader(open(tmp_path / "t.csv", encoding="utf-8")))
    assert rows == [["record", "r", "v"], ["a", "1.5", "2.0"], ["b", "", "3.0"]]


def test_formulas_columns_and_fit_and_only_arithmetic_is_allowed():
    assert names_in("a*x**b + c") == ["a", "x", "b", "c"] and parameters_of("a*x**b + c") == ["a", "b", "c"]
    assert evaluate("sqrt(vx**2 + vy**2)", {"vx": np.array([3.0]), "vy": np.array([4.0])}).tolist() == [5.0]
    for bad in ('__import__("os")', "x.real", "open(1)", '"text"', "(lambda: 1)()"):
        with pytest.raises(ValueError):
            evaluate(bad, {"x": 1.0})  # формулой ничего, кроме вычислений, не сделать
    table = DataTable({"F": np.array([140.0, 160.0]), "record": np.array(["a", "b"], dtype=object)})
    extended, errors = computed_columns(table, "U = 0.07*F\nU2 = 2*U  # колонка из колонки выше\nплохо\nF = 1\nW = Q")
    assert extended.columns["U2"].tolist() == [19.6, 22.400000000000002] and len(errors) == 3
    x = np.linspace(1.0, 10.0, 25)
    result = fit("a*x**b + c", x, 2.5 * x**1.7 + 3.0, {"a": 1.0, "b": 1.0, "c": 0.0})
    assert result.ok and np.allclose([result.values[p] for p in "abc"], [2.5, 1.7, 3.0], rtol=1e-4) and result.r2 > 0.999
    assert not fit("a*x", x[:0], x[:0], {}).ok  # нет точек — понятный отказ, а не ошибка
