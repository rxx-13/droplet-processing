"""Таблица записей (xlsx), выбор записей и общий цикл обработки стадии.

Таблица: первый лист; строка 1 — заголовки, строка 2 — описания (пропускается), данные с 3-й строки.
Колонки:
    rec_name, images_path_side        — обязательные;
    cond_num                          — номер условий (для выбора записей, CONDITIONS в settings);
    start_frame, fin_frame            — диапазон кадров (номера файлов, с 1), нужен st1;
    rim или rim_new                   — порог площади рима и крупных брызг, px² (st1);
    cx или c_x, cy или c_y            — начало координат для st4, px;
    bag_hight, bag_length, bag_frame  — высота и длина бэга, мм, и кадр, на котором они измерены
                                        (заполняет приложение; стадиям не нужны).
Строки без rec_name или images_path_side пропускаются.

Путь записи: images_path_side, а если он не оканчивается именем записи — images_path_side\\rec_name.
От него строятся папка кадров и выходные папки (имена как в MATLAB):
    <IMAGES_DIR>/<путь записи>/                              кадры
    <OUTPUT_DIR>/PIV/<путь записи>_output_st1_droplets/      результат st1
    <OUTPUT_DIR>/PTV/<путь записи>_output_st2_PTV/           результат st2
    <OUTPUT_DIR>/st4/<путь записи>_output_st4/               результат st4
Пути в settings.py (IMAGES_DIR, TABLE_FILE, OUTPUT_DIR, MATLAB_RESULTS_DIR) — полные или от папки проекта,
например experiments\\example1\\results (project_path): так примеры работают, где бы ни лежал проект.
"""

import logging
import math
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from types import ModuleType
from typing import Any

from openpyxl import load_workbook

log = logging.getLogger(__name__)

PROJECT_DIR = Path(__file__).resolve().parents[2]  # папка проекта: от неё — относительные пути в settings.py

# Названия колонок: в разных таблицах одна и та же величина называется по-разному.
COLUMN_NAMES = {
    "cond_num": ("cond_num",),
    "rec_name": ("rec_name",),
    "images_path_side": ("images_path_side",),
    "start_frame": ("start_frame",),
    "fin_frame": ("fin_frame",),
    "rim": ("rim_new", "rim"),
    "cx": ("cx", "c_x"),
    "cy": ("cy", "c_y"),
    "bag_height": ("bag_hight", "bag_height"),  # размер бэга (приложение, стадиям не нужен)
    "bag_length": ("bag_length",),
    "bag_frame": ("bag_frame",),
}
REQUIRED_COLUMNS = ("rec_name", "images_path_side")


def project_path(path: str | Path) -> Path:
    """Путь из settings.py: полный — как есть, относительный — от папки проекта."""
    path = Path(path)
    return path if path.is_absolute() else PROJECT_DIR / path


@dataclass(frozen=True)
class Record:
    """Одна строка таблицы записей.

    name — rec_name; path — путь записи относительно IMAGES_DIR (разделитель '/'); cond_num — номер
    условий; start_frame, fin_frame — диапазон кадров (номера файлов, с 1); rim — порог площади рима, px²;
    cx, cy — начало координат st4, px. Отсутствующее в таблице — None.
    """

    name: str
    path: str  # путь записи относительно IMAGES_DIR, разделитель '/'
    cond_num: float | None
    start_frame: int | None
    fin_frame: int | None
    rim: float | None
    cx: float | None
    cy: float | None

    @property
    def has_frames(self) -> bool:
        """Задан ли правильный диапазон кадров (без него запись не обрабатывается)."""
        return (self.start_frame is not None and self.fin_frame is not None
                and self.fin_frame >= self.start_frame >= 1)

    def images_dir(self, images_root: Path) -> Path:
        """Папка кадров: <IMAGES_DIR>/<путь записи> (IMAGES_DIR — полный или от папки проекта)."""
        return project_path(images_root) / self.path

    def st1_dir(self, output_root: Path) -> Path:
        """Папка результата st1: <OUTPUT_DIR>/PIV/<путь записи>_output_st1_droplets."""
        return project_path(output_root) / "PIV" / f"{self.path}_output_st1_droplets"

    def st2_dir(self, output_root: Path) -> Path:
        """Папка результата st2: <OUTPUT_DIR>/PTV/<путь записи>_output_st2_PTV."""
        return project_path(output_root) / "PTV" / f"{self.path}_output_st2_PTV"

    def st4_dir(self, output_root: Path) -> Path:
        """Папка результата st4: <OUTPUT_DIR>/st4/<путь записи>_output_st4."""
        return project_path(output_root) / "st4" / f"{self.path}_output_st4"


def _text(value: Any) -> str:
    """Значение ячейки как текст (числа без «.0»: имя записи может быть числом)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value: Any) -> float | None:
    """Значение ячейки как число; пустая ячейка, текст или NaN — None."""
    if value is None or isinstance(value, str):
        return None
    f = float(value)
    return None if math.isnan(f) else f


def _integer(value: Any) -> int | None:
    """Значение ячейки как целое (округление); пустое — None."""
    f = _number(value)
    return None if f is None else int(round(f))


def read_records(table_file: Path, conditions: list[float] | tuple[float, ...] = (),
                 names: list[str] | tuple[str, ...] = ()) -> list[Record]:
    """Записи из таблицы table_file (полный путь или от папки проекта); conditions — номера условий,
    names — имена записей (пусто = все)."""
    table_file = project_path(table_file)
    workbook = load_workbook(table_file, data_only=True, read_only=True)
    rows = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    workbook.close()

    headers = [_text(h) for h in rows[0]]
    column: dict[str, int | None] = {}
    for key, variants in COLUMN_NAMES.items():
        column[key] = next((headers.index(v) for v in variants if v in headers), None)
        if column[key] is None and key in REQUIRED_COLUMNS:
            raise KeyError(f"В таблице {table_file.name} нет колонки {' / '.join(variants)}")

    def cell(row: list[Any], key: str) -> Any:
        i = column[key]
        return row[i] if i is not None and i < len(row) else None

    records = []
    for row in rows[2:]:
        name, side = _text(cell(row, "rec_name")), _text(cell(row, "images_path_side"))
        if not name or not side:
            continue
        path = side.replace("\\", "/").rstrip("/")
        if not path.endswith("/" + name) and path != name:
            path = f"{path}/{name}"
        records.append(Record(
            name=name, path=path, cond_num=_number(cell(row, "cond_num")),
            start_frame=_integer(cell(row, "start_frame")), fin_frame=_integer(cell(row, "fin_frame")),
            rim=_number(cell(row, "rim")), cx=_number(cell(row, "cx")), cy=_number(cell(row, "cy")),
        ))
    log.info("Таблица %s: %d записей", table_file.name, len(records))

    if conditions:
        records = [r for r in records if r.cond_num in conditions]
    if names:
        missing = set(names) - {r.name for r in records}
        if missing:
            log.warning("Нет в таблице (или не подходят под CONDITIONS): %s", ", ".join(sorted(missing)))
        records = [r for r in records if r.name in names]
    return records


def run_stage(stage: str, settings: ModuleType, thresholds: ModuleType,
              process_record: Callable[[Record, ModuleType, ModuleType], None]) -> int:
    """Общий цикл стадии: выбирает записи по settings и обрабатывает каждую.

    Ошибка в одной записи не останавливает остальные. Записи без start_frame/fin_frame
    пропускаются. Возвращает код выхода: 0 — всё обработано, 1 — были ошибки.
    """
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)-7s %(message)s",
                        datefmt="%H:%M:%S")
    records = read_records(Path(settings.TABLE_FILE), settings.CONDITIONS, settings.RECORDS)
    skipped = [r.name for r in records if not r.has_frames]
    if skipped:
        log.warning("Пропущены записи без start_frame/fin_frame (%d): %s", len(skipped), ", ".join(skipped))
    records = [r for r in records if r.has_frames]
    if not records:
        log.error("Нет записей для обработки: проверьте CONDITIONS и RECORDS в settings")
        return 1
    done, failed = 0, 0
    for i, record in enumerate(records, start=1):
        log.info("%s: запись %s (%d из %d)", stage, record.name, i, len(records))
        try:
            process_record(record, settings, thresholds)
            done += 1
        except FileNotFoundError as error:
            log.error("%s: %s", record.name, error)
            failed += 1
        except Exception:  # одна запись не должна останавливать весь прогон
            log.exception("%s: ошибка обработки", record.name)
            failed += 1
    log.info("%s: готово %d, с ошибкой %d, пропущено %d", stage, done, failed, len(skipped))
    return 1 if failed else 0
