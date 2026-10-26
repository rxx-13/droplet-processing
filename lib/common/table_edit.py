"""Таблица записей (xlsx) для приложения: все колонки, несохранённые правки, автозаполнение, запись.

Чтение — по тем же правилам, что read_records (lib/common/records.py): первый лист, строка 1 — заголовки,
строка 2 — описания колонок, данные — с 3-й строки. Но здесь берутся все колонки (их набор — из заголовков,
ничего не зашито) и все строки, где есть rec_name.

Правки копятся в памяти (EditableTable.set) и попадают в файл только по EditableTable.save: книга
открывается заново, меняются только изменённые ячейки, остальное (другие ячейки, форматирование, условное
форматирование, ширина колонок) остаётся как было. Перед записью проверяется, что в строке та же запись
(rec_name) — на случай, если таблицу тем временем правили в Excel. Если файл открыт в Excel, он не
трогается, а возникает TableSaveError с понятным текстом.

Автозаполнение: подсказка для пустого поля — значение той же колонки у предыдущей заполненной записи
(заполненная — с диапазоном кадров); «как у предыдущих» — колонки, одинаковые во всех заполненных записях.
"""

import io
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .records import COLUMN_NAMES, REQUIRED_COLUMNS, Record, _integer, _number, _text


class TableSaveError(Exception):
    """Таблицу не удалось записать (открыта в Excel, строки переставлены и т. п.); файл не изменён."""


@dataclass
class TableRow:
    """Строка данных таблицы: excel_row — номер строки в Excel (с 1); values — значения по заголовкам."""

    excel_row: int
    values: dict[str, Any]


def is_empty(value: Any) -> bool:
    """Пустая ли ячейка: None или пустой текст."""
    return value is None or (isinstance(value, str) and not value.strip())


def same_value(a: Any, b: Any) -> bool:
    """Одинаковы ли значения ячеек: пустые равны пустым, числа сравниваются как числа (200 = 200.0)."""
    if is_empty(a) or is_empty(b):
        return is_empty(a) and is_empty(b)
    numbers = (int, float)
    if isinstance(a, numbers) and isinstance(b, numbers) and not isinstance(a, bool) and not isinstance(b, bool):
        return float(a) == float(b)
    return type(a) is type(b) and a == b


def format_value(value: Any) -> str:
    """Значение ячейки как текст для поля ввода (целые числа — без «.0»)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_table(path: Path) -> tuple[list[str], dict[str, str], list[TableRow]]:
    """Заголовки (непустые, по порядку), описания колонок (строка 2) и строки данных, где есть rec_name."""
    workbook = load_workbook(path, data_only=True)
    rows = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    workbook.close()
    columns: list[tuple[int, str]] = []  # (номер колонки с 0, заголовок)
    for i, header in enumerate(_text(h) for h in (rows[0] if rows else [])):
        if header and header not in {h for _, h in columns}:
            columns.append((i, header))
    headers = [h for _, h in columns]
    name = next((h for h in COLUMN_NAMES["rec_name"] if h in headers), None)
    if name is None:
        raise KeyError(f"В таблице {path.name} нет колонки rec_name")

    def cell(row: list[Any], i: int) -> Any:
        return row[i] if i < len(row) else None

    description_row = rows[1] if len(rows) > 1 else []
    descriptions = {h: _text(cell(description_row, i)) for i, h in columns}
    data = []
    for excel_row, row in enumerate(rows[2:], start=3):
        values = {h: cell(row, i) for i, h in columns}
        if _text(values[name]):
            data.append(TableRow(excel_row, values))
    return headers, descriptions, data


def save_cells(path: Path, cells: dict[tuple[int, str], Any], names: dict[int, str], name_header: str) -> None:
    """Записать ячейки {(строка Excel, заголовок): значение} в первый лист файла, остальное не трогая.

    names — {строка Excel: rec_name, который должен в ней стоять}; не совпало — TableSaveError.
    Книга сначала собирается в памяти и только потом пишется в файл одним разом: если файл открыт
    в Excel, он остаётся прежним.
    """
    workbook = load_workbook(path)
    sheet = workbook.worksheets[0]
    headers = [_text(c.value) for c in sheet[1]]
    for header in {name_header, *(h for _, h in cells)}:
        if header not in headers:
            raise TableSaveError(f"В таблице {path.name} больше нет колонки {header}")
    name_column = headers.index(name_header) + 1
    for row, expected in names.items():
        actual = _text(sheet.cell(row, name_column).value)
        if actual != expected:
            raise TableSaveError(
                f"В строке {row} таблицы {path.name} теперь «{actual}», а не «{expected}» — таблицу правили "
                "в Excel? Перечитайте её (F5) и повторите правку.")
    for (row, header), value in cells.items():
        sheet.cell(row, headers.index(header) + 1).value = value
    buffer = io.BytesIO()
    workbook.save(buffer)
    try:
        path.write_bytes(buffer.getvalue())
    except PermissionError:
        raise TableSaveError(
            f"Не удалось записать {path.name}: таблица открыта в Excel. Закройте её в Excel "
            "и сохраните ещё раз.") from None


class EditableTable:
    """Таблица записей с несохранёнными правками.

    Строки — по номеру i (с 0, в порядке таблицы), колонки — по заголовку. value(i, h) — значение с учётом
    правок; set(i, h, v) — правка (значение, равное исходному, правку снимает); save — запись в файл.
    path — файл таблицы (TABLE_FILE); headers, descriptions, rows — как возвращает read_table.
    """

    def __init__(self, path: Path) -> None:
        self.path = Path(path)
        self.headers, self.descriptions, self.rows = read_table(self.path)
        self.changes: dict[tuple[int, str], Any] = {}

    def __len__(self) -> int:
        return len(self.rows)

    # --- значения
    def column(self, key: str) -> str | None:
        """Заголовок колонки величины key из COLUMN_NAMES (например, "rim" -> "rim" или "rim_new")."""
        return next((h for h in COLUMN_NAMES[key] if h in self.headers), None)

    def original(self, i: int, header: str) -> Any:
        """Значение ячейки в файле (без правок)."""
        return self.rows[i].values.get(header)

    def value(self, i: int, header: str) -> Any:
        """Значение ячейки с учётом правок."""
        return self.changes.get((i, header), self.original(i, header))

    def set(self, i: int, header: str, value: Any) -> None:
        """Правка ячейки; значение, равное исходному, снимает правку."""
        if same_value(value, self.original(i, header)):
            self.changes.pop((i, header), None)
        else:
            self.changes[(i, header)] = value

    def changed(self, i: int, header: str | None = None) -> bool:
        """Есть ли несохранённая правка в строке i (или в её ячейке header)."""
        if header is not None:
            return (i, header) in self.changes
        return any(row == i for row, _ in self.changes)

    def name(self, i: int) -> str:
        """rec_name строки i (с учётом правок)."""
        return _text(self.value(i, self.column("rec_name")))

    # --- ввод
    def is_text(self, header: str) -> bool:
        """Текстовая ли колонка: rec_name, images_path_side или колонка, где есть только текст."""
        if header in {self.column(key) for key in REQUIRED_COLUMNS}:
            return True
        values = [r.values.get(header) for r in self.rows if not is_empty(r.values.get(header))]
        return bool(values) and all(isinstance(v, str) for v in values)

    def is_numeric(self, header: str) -> bool:
        """Числовая ли колонка: в ней есть числа и нет текста."""
        values = [r.values.get(header) for r in self.rows if not is_empty(r.values.get(header))]
        return bool(values) and all(isinstance(v, (int, float)) for v in values)

    def parse(self, header: str, text: str) -> Any:
        """Текст из поля ввода -> значение ячейки: пусто -> None, в нетекстовой колонке число -> число
        (запятая тоже годится как десятичный знак). Не число в числовой колонке — ValueError."""
        text = text.strip()
        if not text:
            return None
        if self.is_text(header):
            return text
        try:
            number = float(text.replace(",", "."))
        except ValueError:
            if self.is_numeric(header):
                raise ValueError(f"в колонке {header} должно быть число") from None
            return text
        return int(number) if number.is_integer() and abs(number) < 1e15 else number

    # --- запись таблицы как Record
    def record(self, i: int) -> Record | None:
        """Строка i как Record (с учётом правок, по тем же правилам, что read_records);
        None — нет rec_name или images_path_side."""

        def get(key: str) -> Any:
            header = self.column(key)
            return None if header is None else self.value(i, header)

        name, side = _text(get("rec_name")), _text(get("images_path_side"))
        if not name or not side:
            return None
        path = side.replace("\\", "/").rstrip("/")
        if not path.endswith("/" + name) and path != name:  # как в read_records
            path = f"{path}/{name}"
        return Record(name=name, path=path, cond_num=_number(get("cond_num")),
                      start_frame=_integer(get("start_frame")), fin_frame=_integer(get("fin_frame")),
                      rim=_number(get("rim")), cx=_number(get("cx")), cy=_number(get("cy")))

    # --- автозаполнение
    def filled(self, i: int) -> bool:
        """Заполнена ли запись: задан диапазон кадров."""
        record = self.record(i)
        return record is not None and record.has_frames

    def previous_filled(self, i: int) -> int | None:
        """Ближайшая выше заполненная запись; если выше нет — ближайшая ниже; нет совсем — None."""
        above = next((j for j in range(i - 1, -1, -1) if self.filled(j)), None)
        return above if above is not None else next((j for j in range(i + 1, len(self)) if self.filled(j)), None)

    def suggestion(self, i: int, header: str) -> tuple[int, Any] | None:
        """Подсказка для пустой ячейки: (строка предыдущей заполненной записи, её значение); иначе None."""
        if not is_empty(self.value(i, header)):
            return None
        j = self.previous_filled(i)
        if j is None or is_empty(self.value(j, header)):
            return None
        return j, self.value(j, header)

    def common_values(self, exclude: int | None = None) -> dict[str, Any]:
        """Колонки, значения которых одинаковы во всех заполненных записях (кроме строки exclude);
        заполненных записей должно быть хотя бы две."""
        filled = [j for j in range(len(self)) if j != exclude and self.filled(j)]
        if len(filled) < 2:
            return {}
        common = {}
        for header in self.headers:
            values = [self.value(j, header) for j in filled]
            if not is_empty(values[0]) and all(same_value(v, values[0]) for v in values):
                common[header] = values[0]
        return common

    def fill_like_previous(self, i: int) -> list[str]:
        """Заполнить пустые ячейки строки i значениями, одинаковыми во всех заполненных записях;
        возвращает заголовки заполненных колонок."""
        filled = []
        for header, value in self.common_values(exclude=i).items():
            if is_empty(self.value(i, header)):
                self.set(i, header, value)
                filled.append(header)
        return filled

    # --- файл
    def save(self, backup_dir: Path | None = None) -> Path | None:
        """Записать правки в файл (только изменённые ячейки) и перечитать таблицу.

        backup_dir — папка, куда сначала кладётся копия файла (None — без копии). Возвращает путь копии
        (или None). Не удалось записать — TableSaveError, правки остаются несохранёнными.
        """
        if not self.changes:
            return None
        backup = None
        if backup_dir is not None:
            backup_dir.mkdir(parents=True, exist_ok=True)
            backup = backup_dir / f"{self.path.stem} {datetime.now():%Y-%m-%d %H-%M-%S}{self.path.suffix}"
            shutil.copy2(self.path, backup)
        name_header = self.column("rec_name")
        cells = {(self.rows[i].excel_row, header): value for (i, header), value in self.changes.items()}
        names = {self.rows[i].excel_row: _text(self.original(i, name_header)) for i, _ in self.changes}
        save_cells(self.path, cells, names, name_header)
        self.headers, self.descriptions, self.rows = read_table(self.path)
        self.changes = {}
        return backup
